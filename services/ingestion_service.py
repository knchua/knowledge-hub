"""Extract page-aware source text and split with LangChain."""
from __future__ import annotations
import io
from dataclasses import dataclass

from services.knowledge_service import KnowledgeError, validate_file

CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
MAX_CHUNKS = 1500

@dataclass(frozen=True)
class TextChunk:
    index: int
    content: str
    page_number: int | None
    source_metadata: dict


def extract_chunks(filename: str, file_bytes: bytes) -> list[TextChunk]:
    validated = validate_file(filename, file_bytes)
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    from langchain_core.documents import Document
    splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
    sources: list[Document] = []
    try:
        if validated.extension == ".pdf":
            import fitz  # PyMuPDF
            with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
                if pdf.needs_pass:
                    raise KnowledgeError("Password-protected PDFs are not supported.")
                for idx, page in enumerate(pdf):
                    text = page.get_text("text")
                    if text.strip():
                        sources.append(Document(page_content=text, metadata={"page_number": idx + 1}))
        elif validated.extension == ".docx":
            from docx import Document as WordDocument
            doc = WordDocument(io.BytesIO(file_bytes))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    line = " | ".join(cell.text.strip() for cell in row.cells)
                    if line.strip(" |"):
                        parts.append(line)
            text = "\n".join(parts)
            if text.strip():
                sources.append(Document(page_content=text, metadata={}))
        else:
            text = file_bytes.decode("utf-8-sig")
            if text.strip():
                sources.append(Document(page_content=text, metadata={}))
    except KnowledgeError:
        raise
    except Exception as exc:
        raise KnowledgeError("Document text could not be extracted. Check the file format.") from exc
    if not sources:
        raise KnowledgeError("No readable text was found; scanned PDFs require OCR (out of scope).")
    chunks: list[TextChunk] = []
    for source in sources:
        for section in splitter.split_documents([source]):
            value = section.page_content.strip()
            if value:
                index = len(chunks)
                chunks.append(TextChunk(index, value, section.metadata.get("page_number"), {
                    "source_filename": validated.display_name,
                    "page_number": section.metadata.get("page_number"),
                    "chunk_index": index,
                }))
                if len(chunks) > MAX_CHUNKS:
                    raise KnowledgeError("Document exceeds the maximum 1,500 text chunks.")
    if not chunks:
        raise KnowledgeError("No readable text chunks were extracted.")
    return chunks
